"""Render the Lean Master workbook as the repository requirements register."""

from __future__ import annotations

import argparse
import re
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = (
    "ID",
    "Stage",
    "Requirement",
    "Simple check",
    "Next choice",
    "Student owner",
    "Status",
    "Group",
    "Basis",
    "Previous master IDs",
)
TRANSLATION = str.maketrans(
    {
        "°": " degrees ",
        "–": "-",
        "—": "-",
        "’": "'",
        "‘": "'",
        "“": '"',
        "”": '"',
        "≥": ">=",
        "≤": "<=",
        "→": " to ",
    }
)


def clean(value: Any) -> str:
    """Make a cell safe and deterministic for a Markdown table cell."""
    text = "" if value is None else " ".join(str(value).split())
    return text.translate(TRANSLATION).replace("|", r"\|")


def find_header_row(worksheet: Any) -> dict[str, int]:
    for row in worksheet.iter_rows(min_row=1, max_row=min(25, worksheet.max_row), values_only=True):
        labels = {clean(value): index for index, value in enumerate(row)}
        if all(column in labels for column in REQUIRED_COLUMNS):
            return labels
    expected = ", ".join(REQUIRED_COLUMNS)
    raise ValueError(f"Master sheet is missing the required header row: {expected}")


def workbook_rows(workbook_path: Path) -> tuple[OrderedDict[str, list[dict[str, str]]], str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return rows_from_workbook(workbook)
    finally:
        workbook.close()


def rows_from_workbook(workbook: Any) -> tuple[OrderedDict[str, list[dict[str, str]]], str]:
    if "Master" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a sheet named 'Master'.")

    worksheet = workbook["Master"]
    header = find_header_row(worksheet)
    groups: OrderedDict[str, list[dict[str, str]]] = OrderedDict()
    header_row = next(
        row_index
        for row_index in range(1, min(25, worksheet.max_row) + 1)
        if all(
            clean(worksheet.cell(row_index, header[column] + 1).value) == column
            for column in REQUIRED_COLUMNS
        )
    )

    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = {column: clean(row[index]) if index < len(row) else "" for column, index in header.items()}
        requirement_id = values["ID"]
        if not requirement_id:
            continue
        if not re.fullmatch(r"R\d+", requirement_id):
            raise ValueError(f"Invalid requirement ID: {requirement_id}")
        if not values["Requirement"] or not values["Simple check"] or not values["Group"]:
            raise ValueError(f"{requirement_id} is missing a requirement, simple check, or group.")
        groups.setdefault(values["Group"], []).append(values)

    if not groups:
        raise ValueError("Master sheet contains no requirements.")

    metadata = clean(worksheet["E3"].value)
    return groups, metadata


def render(workbook_path: Path) -> str:
    groups, metadata = workbook_rows(workbook_path)
    metadata_sentence = f" The workbook metadata says: {metadata}" if metadata else ""
    lines = [
        "# Master requirements",
        "",
        "This register is generated from `URC_2027_Lean_Master.xlsx`. The SharePoint workbook is the source of truth; the committed workbook and this Markdown register are GitHub mirrors." + metadata_sentence,
        "",
        "The workbook labels all entries as proposed and uses 2026 references only. Confirm every competition-derived item against the current 2027 rulebook before approving a design.",
        "",
        "Each entry is a team requirement rather than a replacement for the rulebook. Source / basis preserves the workbook's basis label and prior-master traceability. Status combines the workbook status and delivery stage.",
        "",
        "## Requirement format",
        "",
        "| ID | Requirement | Source / basis | Verification | Status |",
        "| --- | --- | --- | --- | --- |",
        "| REQ-XXX-001 | _Write a testable shall statement._ | Rulebook section X.Y or documented team basis | Inspection, analysis, test, or demonstration | Draft - Build first |",
    ]

    for group, entries in groups.items():
        lines.extend(("", f"## {group}", "", "| ID | Requirement | Source / basis | Verification | Status |", "| --- | --- | --- | --- | --- |"))
        for entry in entries:
            source = entry["Basis"]
            if entry["Previous master IDs"]:
                source = f"{source}; prior IDs: {entry['Previous master IDs']}"
            status = f"{entry['Status']} - {entry['Stage']}"
            lines.append(
                f"| {entry['ID']} | {entry['Requirement']} | {source} | "
                f"{entry['Simple check']} | {status} |"
            )

        choices = [entry for entry in entries if entry["Next choice"]]
        if choices:
            lines.extend(("", "### Open choices and implementation notes", "", "| Requirement | Open choice or note |", "| --- | --- |"))
            lines.extend(f"| {entry['ID']} | {entry['Next choice']} |" for entry in choices)

    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--check", action="store_true", help="Fail if the output is not current.")
    args = parser.parse_args()

    rendered = render(args.workbook)
    current = args.output.read_text(encoding="utf-8") if args.output.exists() else ""
    if current == rendered:
        return 0
    if args.check:
        print(f"{args.output} is not current with {args.workbook}.", file=sys.stderr)
        return 1

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(rendered, encoding="utf-8", newline="\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
